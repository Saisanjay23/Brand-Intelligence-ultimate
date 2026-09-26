"""Pins for the defects found in the 2026-09-25 codebase audit.

Each class is one defect that was live in the code, with what it did:

  * a download header that crashed on a non-Latin client name;
  * an .xlsx export that wrote a visible apostrophe into cells to stop
    formulas, when the cell type is what stops them;
  * an SMTP path that fell through to a PLAINTEXT login after STARTTLS
    failed, and alert emails that did not escape what they interpolated;
  * "analyse these ids" skipping the per-call ceiling the other path has;
  * a scheduler queue that would sweep a repeated client twice;
  * fire-and-forget tasks that asyncio could garbage-collect mid-flight.

PURE LOGIC ONLY, like the rest of this suite.
"""

from __future__ import annotations

import asyncio
import gc
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch

import pytest


class TestDownloadHeaders:
    def test_a_non_latin_name_is_a_valid_header(self):
        """HTTP headers are Latin-1. This value used to go straight in, so
        a client named in Devanagari turned its report into a 500."""
        from backend.api.models import content_disposition

        v = content_disposition("attachment", "report-अदाणी-2026-09-25.html", "report.html")
        v.encode("latin-1")                                  # must not raise
        assert 'filename="report--2026-09-25.html"' in v
        assert "filename*=UTF-8''report-%E0%A4%85" in v

    def test_a_plain_name_is_left_alone(self):
        from backend.api.models import content_disposition

        assert content_disposition("attachment", "analysis.xlsx") == \
            'attachment; filename="analysis.xlsx"'

    def test_quotes_and_newlines_cannot_break_out(self):
        from backend.api.models import content_disposition

        v = content_disposition("inline", 'a"b\r\nX-Evil: 1.png', "x.png")
        assert "\r" not in v and "\n" not in v
        assert v.count('"') == 2

    def test_nothing_usable_falls_back(self):
        from backend.api.models import content_disposition

        v = content_disposition("attachment", "अदाणी", "report.html")
        assert v.startswith('attachment; filename="report.html"')


class TestXlsxExport:
    def _export(self, rows, filename="analysis.xlsx"):
        from openpyxl import load_workbook

        from backend.api.analysis import ExportXlsx, export_xlsx

        resp = asyncio.run(export_xlsx(ExportXlsx(filename=filename, rows=rows)))
        ws = load_workbook(BytesIO(resp.body)).active
        return resp, [[c.value for c in r] for r in ws.iter_rows()], ws

    def test_a_formula_is_stored_as_text_not_run(self):
        resp, cells, ws = self._export([{"name": '=HYPERLINK("http://x","y")'}])
        assert cells[1][0] == '=HYPERLINK("http://x","y")'
        assert ws["A2"].data_type == "s"

    def test_text_is_not_given_a_visible_apostrophe(self):
        """The old guard wrote "'-Official-" into the cell."""
        _, cells, _ = self._export([{"name": "-Official-", "handle": "@brand", "n": "12"}])
        assert cells[1] == ["-Official-", "@brand", 12]

    def test_a_non_latin_filename_exports(self):
        resp, _, _ = self._export([{"a": 1}], filename="अदाणी.xlsx")
        resp.headers["content-disposition"].encode("latin-1")


class TestAlertEmails:
    def test_values_are_escaped(self):
        from backend.services import email_service

        out = email_service._format_html_template(
            title="t", badge_text="b", badge_color="#fff", headline="<b>x</b>",
            details_table=[("Account", "<script>alert(1)</script> & co")],
            recommendation="ok",
        )
        assert "<script>" not in out
        assert "&lt;script&gt;alert(1)&lt;/script&gt; &amp; co" in out
        assert "&lt;b&gt;x&lt;/b&gt;" in out

    @pytest.mark.asyncio
    async def test_every_sender_renders(self):
        """Each sender names its result `html`, which shadows the module
        of the same name -- an escape call written as `html.escape` inside
        one of them raises UnboundLocalError. All three must get as far as
        handing a message to send_email."""
        from backend.services import email_service as E

        settings = AsyncMock(return_value={})
        send = AsyncMock(return_value=(True, "ok"))
        with patch.object(E.alert_settings_db, "get_settings", settings), \
             patch.object(E, "send_email", send):
            assert await E.send_session_failure_alert("facebook", "<me>", "s1", "expired")
            assert await E.send_session_expiring_alert("twitter", "a&b", 3.0)
            assert await E.send_critical_incident_alert({"fix": "<do this>"})
        body = send.await_args_list[0].args[1]
        assert "&lt;me&gt;" in body

    def test_starttls_failure_never_logs_in_in_plaintext(self):
        from email.mime.multipart import MIMEMultipart

        from backend.services import email_service as E

        server = MagicMock()
        server.has_extn.return_value = True
        server.starttls.side_effect = OSError("handshake failed")
        smtp = MagicMock()
        smtp.return_value.__enter__.return_value = server
        with patch.object(E.smtplib, "SMTP", smtp):
            ok, detail = E._send_smtp_sync(
                "smtp.example.com", 587, "user", "secret", "a@b", ["c@d"], MIMEMultipart())
        assert ok is False
        assert "STARTTLS" in detail
        server.login.assert_not_called()
        server.sendmail.assert_not_called()


class TestAnalyseSelectedIds:
    @pytest.mark.asyncio
    async def test_the_ceiling_applies_and_the_overflow_is_named(self, monkeypatch):
        from backend.api import discovery as D

        cap = D._MAX_VALIDATED_PER_ANALYSE
        docs = [{"id": str(i), "url": f"https://x.com/{i}", "status": "approved"}
                for i in range(cap + 7)]
        monkeypatch.setattr(D.profiles_db, "get_by_ids", AsyncMock(return_value=docs))
        job = MagicMock(id="j1", status="queued", total=cap)
        start = AsyncMock(return_value=(job, []))
        monkeypatch.setattr(D.analysis_runner, "start", start)

        res = await D.analyse_validated(D.AnalyseValidated(
            group_id="g", ids=[d["id"] for d in docs]))
        assert len(start.await_args.args[0]) == cap
        over = [s for s in res.skipped if s.value == "(over limit)"]
        assert over and "7 profile(s)" in over[0].reason

    @pytest.mark.asyncio
    async def test_a_repeated_id_is_not_reported_missing(self, monkeypatch):
        from backend.api import discovery as D

        doc = {"id": "a", "url": "https://x.com/a", "status": "approved"}
        get = AsyncMock(return_value=[doc])
        monkeypatch.setattr(D.profiles_db, "get_by_ids", get)
        job = MagicMock(id="j1", status="queued", total=1)
        monkeypatch.setattr(D.analysis_runner, "start", AsyncMock(return_value=(job, [])))

        res = await D.analyse_validated(D.AnalyseValidated(group_id="g", ids=["a", "a", "a"]))
        assert get.await_args.args[1] == ["a"]
        assert not any(s.value == "(unresolved ids)" for s in res.skipped)


class TestSchedulerQueue:
    @pytest.mark.asyncio
    async def test_a_client_is_queued_once(self, monkeypatch):
        from backend.api import scheduler as S

        engine = MagicMock(busy=False, snapshot=AsyncMock(return_value={"schedule": {}}))
        monkeypatch.setattr(S, "scheduler_engine", engine)
        monkeypatch.setattr(S.clients_db, "try_get", AsyncMock(return_value=None))
        saved = AsyncMock()
        monkeypatch.setattr(S.schedule_db, "set_queue", saved)
        monkeypatch.setattr(S, "_state_out", lambda snap, sched: {})

        await S.set_queue(S.QueueBody(client_ids=["b", "a", "b", " ", "a", "c"]))
        assert saved.await_args.args[0] == ["b", "a", "c"]


class TestSpawnedTasksAreHeld:
    def test_a_spawned_task_survives_garbage_collection(self):
        from backend.shared import tasks

        async def go():
            done = asyncio.Event()

            async def work():
                await asyncio.sleep(0.01)
                done.set()

            tasks.spawn(work())
            gc.collect()
            assert tasks.in_flight() >= 1
            await asyncio.wait_for(done.wait(), 1)
            await asyncio.sleep(0)
            assert tasks.in_flight() == 0

        asyncio.run(go())

    def test_no_bare_response_handlers_are_left(self):
        """Every platform engine's `page.on("response", ...)` used to start
        an unreferenced task per response -- the responses that carry the
        search results."""
        from pathlib import Path

        for f in Path("backend/platforms").rglob("*.py"):
            src = f.read_text(encoding="utf-8")
            assert "lambda r: asyncio.create_task(" not in src, f
